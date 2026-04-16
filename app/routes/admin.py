import os
from io import BytesIO
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, current_app
from flask_login import login_user, logout_user, login_required, current_user
from werkzeug.utils import secure_filename
from app.models import (db, Admin, Product, Category, Brand, ProductImage, ProductDocument,
                        RFQRequest, FeaturedProduct, ShippingInfo, CompanyInfo)
from app.forms import LoginForm, ProductForm, CategoryForm, BrandForm, ShippingForm
from app.utils import save_uploaded_file, generate_slug

admin_bp = Blueprint('admin_bp', __name__)


# ─── Auth ───
@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('admin_bp.dashboard'))
    form = LoginForm()
    if form.validate_on_submit():
        admin = Admin.query.filter_by(username=form.username.data).first()
        if admin and admin.check_password(form.password.data):
            login_user(admin, remember=form.remember.data)
            return redirect(url_for('admin_bp.dashboard'))
        flash('اسم المستخدم أو كلمة المرور غير صحيحة', 'error')
    return render_template('admin/login.html', form=form)


@admin_bp.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('admin_bp.login'))


# ─── Dashboard ───
@admin_bp.route('/')
@login_required
def dashboard():
    total_products = Product.query.count()
    active_products = Product.query.filter_by(is_active=True).count()
    total_orders = RFQRequest.query.count()
    new_orders = RFQRequest.query.filter_by(status='new').count()
    recent_orders = RFQRequest.query.order_by(RFQRequest.created_at.desc()).limit(10).all()
    return render_template('admin/dashboard.html',
                           total_products=total_products,
                           active_products=active_products,
                           total_orders=total_orders,
                           new_orders=new_orders,
                           recent_orders=recent_orders)


# ─── Products CRUD ───
@admin_bp.route('/products')
@login_required
def products_list():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('q', '')
    category_id = request.args.get('category', type=int)
    brand_id = request.args.get('brand', type=int)

    query = Product.query
    if search:
        query = query.filter(db.or_(Product.name_ar.ilike(f'%{search}%'), Product.sku.ilike(f'%{search}%')))
    if category_id:
        query = query.filter_by(category_id=category_id)
    if brand_id:
        query = query.filter_by(brand_id=brand_id)

    products = query.order_by(Product.created_at.desc()).paginate(page=page, per_page=20, error_out=False)
    categories = Category.query.filter_by(parent_id=None).order_by(Category.sort_order).all()
    brands = Brand.query.order_by(Brand.name).all()

    return render_template('admin/products.html', products=products, categories=categories, brands=brands)


@admin_bp.route('/products/add', methods=['GET', 'POST'])
@login_required
def product_add():
    form = ProductForm()
    form.category_id.choices = [(c.id, c.name_ar) for c in Category.query.filter_by(is_active=True).order_by(Category.sort_order).all()]
    form.brand_id.choices = [(b.id, b.name) for b in Brand.query.filter_by(is_active=True).order_by(Brand.name).all()]

    if form.validate_on_submit():
        product = Product(
            name_ar=form.name_ar.data,
            sku=form.sku.data,
            slug=generate_slug(form.name_ar.data),
            category_id=form.category_id.data,
            brand_id=form.brand_id.data,
            short_description_ar=form.short_description_ar.data,
            full_description_ar=form.full_description_ar.data,
            specs_json=form.specs_json.data if form.specs_json.data else None,
            shipping_info=form.shipping_info.data,
            is_featured=form.is_featured.data,
            is_active=form.is_active.data
        )
        db.session.add(product)
        db.session.flush()

        # Handle image uploads
        images = request.files.getlist('images')
        for i, img_file in enumerate(images):
            if img_file and img_file.filename:
                path = save_uploaded_file(img_file, 'images')
                if path:
                    db.session.add(ProductImage(
                        product_id=product.id,
                        image_path=path,
                        is_primary=(i == 0),
                        sort_order=i
                    ))

        # Handle document uploads
        for doc_type in ['datasheet', 'manual', 'brochure']:
            doc_file = request.files.get(f'doc_{doc_type}')
            if doc_file and doc_file.filename:
                path = save_uploaded_file(doc_file, 'documents')
                if path:
                    db.session.add(ProductDocument(
                        product_id=product.id,
                        title=doc_type,
                        file_path=path,
                        doc_type=doc_type
                    ))

        db.session.commit()
        flash('تم إضافة المنتج بنجاح', 'success')
        return redirect(url_for('admin_bp.products_list'))

    return render_template('admin/product_form.html', form=form, is_edit=False)


@admin_bp.route('/products/edit/<int:product_id>', methods=['GET', 'POST'])
@login_required
def product_edit(product_id):
    product = Product.query.get_or_404(product_id)
    form = ProductForm(obj=product)
    form.category_id.choices = [(c.id, c.name_ar) for c in Category.query.filter_by(is_active=True).order_by(Category.sort_order).all()]
    form.brand_id.choices = [(b.id, b.name) for b in Brand.query.filter_by(is_active=True).order_by(Brand.name).all()]

    if form.validate_on_submit():
        product.name_ar = form.name_ar.data
        product.sku = form.sku.data
        product.slug = generate_slug(form.name_ar.data)
        product.category_id = form.category_id.data
        product.brand_id = form.brand_id.data
        product.short_description_ar = form.short_description_ar.data
        product.full_description_ar = form.full_description_ar.data
        product.specs_json = form.specs_json.data if form.specs_json.data else None
        product.shipping_info = form.shipping_info.data
        product.is_featured = form.is_featured.data
        product.is_active = form.is_active.data

        # Handle new image uploads
        images = request.files.getlist('images')
        for i, img_file in enumerate(images):
            if img_file and img_file.filename:
                path = save_uploaded_file(img_file, 'images')
                if path:
                    db.session.add(ProductImage(
                        product_id=product.id,
                        image_path=path,
                        is_primary=False,
                        sort_order=product.images.count() + i
                    ))

        # Handle document uploads
        for doc_type in ['datasheet', 'manual', 'brochure']:
            doc_file = request.files.get(f'doc_{doc_type}')
            if doc_file and doc_file.filename:
                # Remove old doc of same type
                old = ProductDocument.query.filter_by(product_id=product.id, doc_type=doc_type).first()
                if old:
                    db.session.delete(old)
                path = save_uploaded_file(doc_file, 'documents')
                if path:
                    db.session.add(ProductDocument(
                        product_id=product.id,
                        title=doc_type,
                        file_path=path,
                        doc_type=doc_type
                    ))

        db.session.commit()
        flash('تم تعديل المنتج بنجاح', 'success')
        return redirect(url_for('admin_bp.products_list'))

    return render_template('admin/product_form.html', form=form, product=product, is_edit=True)


@admin_bp.route('/products/delete/<int:product_id>', methods=['POST'])
@login_required
def product_delete(product_id):
    product = Product.query.get_or_404(product_id)
    db.session.delete(product)
    db.session.commit()
    flash('تم حذف المنتج', 'success')
    return redirect(url_for('admin_bp.products_list'))


@admin_bp.route('/products/toggle/<int:product_id>', methods=['POST'])
@login_required
def product_toggle(product_id):
    product = Product.query.get_or_404(product_id)
    product.is_active = not product.is_active
    db.session.commit()
    status = 'مفعل' if product.is_active else 'مخفي'
    flash(f'المنتج الآن {status}', 'success')
    return redirect(url_for('admin_bp.products_list'))


@admin_bp.route('/products/image/delete/<int:image_id>', methods=['POST'])
@login_required
def delete_product_image(image_id):
    img = ProductImage.query.get_or_404(image_id)
    product_id = img.product_id
    db.session.delete(img)
    db.session.commit()
    flash('تم حذف الصورة', 'success')
    return redirect(url_for('admin_bp.product_edit', product_id=product_id))


@admin_bp.route('/products/image/primary/<int:image_id>', methods=['POST'])
@login_required
def set_primary_image(image_id):
    img = ProductImage.query.get_or_404(image_id)
    # Unset all primary for this product
    ProductImage.query.filter_by(product_id=img.product_id).update({'is_primary': False})
    img.is_primary = True
    db.session.commit()
    flash('تم تعيين الصورة الرئيسية', 'success')
    return redirect(url_for('admin_bp.product_edit', product_id=img.product_id))


# ─── Categories ───
@admin_bp.route('/categories')
@login_required
def categories_list():
    categories = Category.query.filter_by(parent_id=None).order_by(Category.sort_order).all()
    return render_template('admin/categories.html', categories=categories)


@admin_bp.route('/categories/add', methods=['POST'])
@login_required
def category_add():
    name_ar = request.form.get('name_ar', '').strip()
    name_en = request.form.get('name_en', '').strip()
    parent_id = request.form.get('parent_id', type=int)
    icon = request.form.get('icon', '').strip()

    if name_ar:
        cat = Category(
            name_ar=name_ar,
            name_en=name_en or None,
            slug=generate_slug(name_en or name_ar),
            icon=icon or None,
            parent_id=parent_id if parent_id else None,
            sort_order=Category.query.count()
        )
        db.session.add(cat)
        db.session.commit()
        flash('تم إضافة التصنيف', 'success')
    return redirect(url_for('admin_bp.categories_list'))


@admin_bp.route('/categories/edit/<int:cat_id>', methods=['POST'])
@login_required
def category_edit(cat_id):
    cat = Category.query.get_or_404(cat_id)
    cat.name_ar = request.form.get('name_ar', cat.name_ar)
    cat.name_en = request.form.get('name_en', cat.name_en)
    cat.slug = generate_slug(request.form.get('name_en') or request.form.get('name_ar') or cat.name_ar)
    cat.icon = request.form.get('icon', cat.icon)
    db.session.commit()
    flash('تم تعديل التصنيف', 'success')
    return redirect(url_for('admin_bp.categories_list'))


@admin_bp.route('/categories/delete/<int:cat_id>', methods=['POST'])
@login_required
def category_delete(cat_id):
    cat = Category.query.get_or_404(cat_id)
    if cat.subcategories.count() > 0:
        flash('لا يمكن حذف تصنيف يحتوي على فئات فرعية، احذف الفئات الفرعية أولاً', 'error')
        return redirect(url_for('admin_bp.categories_list'))
    # Delete all products in this category first
    products = Product.query.filter_by(category_id=cat.id).all()
    product_count = len(products)
    for p in products:
        db.session.delete(p)
    db.session.delete(cat)
    db.session.commit()
    if product_count:
        flash(f'تم حذف التصنيف وجميع منتجاته ({product_count} منتج)', 'success')
    else:
        flash('تم حذف التصنيف', 'success')
    return redirect(url_for('admin_bp.categories_list'))


# ─── Brands ───
@admin_bp.route('/brands')
@login_required
def brands_list():
    brands = Brand.query.order_by(Brand.name).all()
    return render_template('admin/brands.html', brands=brands)


@admin_bp.route('/brands/add', methods=['POST'])
@login_required
def brand_add():
    name = request.form.get('name', '').strip()
    description_ar = request.form.get('description_ar', '').strip()
    website_url = request.form.get('website_url', '').strip()

    if name:
        logo_path = None
        logo = request.files.get('logo')
        if logo and logo.filename:
            logo_path = save_uploaded_file(logo, 'images/brands')

        brand = Brand(
            name=name,
            slug=generate_slug(name),
            logo_path=logo_path,
            description_ar=description_ar or None,
            website_url=website_url or None
        )
        db.session.add(brand)
        db.session.commit()
        flash('تم إضافة العلامة التجارية', 'success')
    return redirect(url_for('admin_bp.brands_list'))


@admin_bp.route('/brands/edit/<int:brand_id>', methods=['POST'])
@login_required
def brand_edit(brand_id):
    brand = Brand.query.get_or_404(brand_id)
    brand.name = request.form.get('name', brand.name)
    brand.slug = generate_slug(brand.name)
    brand.description_ar = request.form.get('description_ar', brand.description_ar)
    brand.website_url = request.form.get('website_url', brand.website_url)

    logo = request.files.get('logo')
    if logo and logo.filename:
        path = save_uploaded_file(logo, 'images/brands')
        if path:
            brand.logo_path = path

    db.session.commit()
    flash('تم تعديل العلامة التجارية', 'success')
    return redirect(url_for('admin_bp.brands_list'))


@admin_bp.route('/brands/delete/<int:brand_id>', methods=['POST'])
@login_required
def brand_delete(brand_id):
    brand = Brand.query.get_or_404(brand_id)
    if brand.product_count > 0:
        flash('لا يمكن حذف علامة تجارية تحتوي على منتجات', 'error')
    else:
        db.session.delete(brand)
        db.session.commit()
        flash('تم حذف العلامة التجارية', 'success')
    return redirect(url_for('admin_bp.brands_list'))


# ─── Orders (RFQ) ───
@admin_bp.route('/orders')
@login_required
def orders_list():
    page = request.args.get('page', 1, type=int)
    status = request.args.get('status', '')
    search = request.args.get('q', '')

    query = RFQRequest.query
    if status:
        query = query.filter_by(status=status)
    if search:
        query = query.filter(db.or_(
            RFQRequest.full_name.ilike(f'%{search}%'),
            RFQRequest.email.ilike(f'%{search}%'),
            RFQRequest.ref_number.ilike(f'%{search}%'),
            RFQRequest.company_name.ilike(f'%{search}%')
        ))

    orders = query.order_by(RFQRequest.created_at.desc()).paginate(page=page, per_page=20, error_out=False)
    return render_template('admin/orders.html', orders=orders)


@admin_bp.route('/orders/<int:order_id>')
@login_required
def order_detail(order_id):
    order = RFQRequest.query.get_or_404(order_id)
    products = []
    if order.product_ids_json:
        products = Product.query.filter(Product.id.in_(order.product_ids_json)).all()
    return render_template('admin/order_detail.html', order=order, bulk_products=products)


@admin_bp.route('/orders/<int:order_id>/status', methods=['POST'])
@login_required
def order_update_status(order_id):
    order = RFQRequest.query.get_or_404(order_id)
    new_status = request.form.get('status')
    if new_status in ('new', 'reviewing', 'replied', 'closed'):
        order.status = new_status
        db.session.commit()
        flash('تم تحديث حالة الطلب', 'success')
    return redirect(url_for('admin_bp.order_detail', order_id=order_id))


@admin_bp.route('/orders/<int:order_id>/notes', methods=['POST'])
@login_required
def order_add_notes(order_id):
    order = RFQRequest.query.get_or_404(order_id)
    order.admin_notes = request.form.get('admin_notes', '')
    db.session.commit()
    flash('تم حفظ الملاحظات', 'success')
    return redirect(url_for('admin_bp.order_detail', order_id=order_id))


@admin_bp.route('/orders/export')
@login_required
def orders_export():
    from openpyxl import Workbook

    orders = RFQRequest.query.order_by(RFQRequest.created_at.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = 'الطلبات'
    ws.sheet_view.rightToLeft = True

    headers = ['رقم المرجع', 'الاسم', 'البريد الإلكتروني', 'الجوال', 'الشركة', 'المنتج', 'الحالة', 'التاريخ', 'ملاحظات الأدمن']
    ws.append(headers)

    for order in orders:
        product_name = order.product.name_ar if order.product else 'طلب متعدد'
        ws.append([
            order.ref_number,
            order.full_name,
            order.email,
            order.phone,
            order.company_name,
            product_name,
            order.status_ar,
            order.created_at.strftime('%Y-%m-%d %H:%M'),
            order.admin_notes or ''
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name='sama_tech_orders.xlsx', as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─── Admins Management ───
@admin_bp.route('/admins')
@login_required
def admins_list():
    admins = Admin.query.order_by(Admin.created_at.desc()).all()
    return render_template('admin/admins.html', admins=admins)


@admin_bp.route('/admins/add', methods=['POST'])
@login_required
def admin_add():
    username = request.form.get('username', '').strip()
    email = request.form.get('email', '').strip()
    password = request.form.get('password', '').strip()

    if not username or not email or not password:
        flash('جميع الحقول مطلوبة', 'error')
        return redirect(url_for('admin_bp.admins_list'))

    if Admin.query.filter_by(username=username).first():
        flash('اسم المستخدم مستخدم بالفعل', 'error')
        return redirect(url_for('admin_bp.admins_list'))

    if Admin.query.filter_by(email=email).first():
        flash('البريد الإلكتروني مستخدم بالفعل', 'error')
        return redirect(url_for('admin_bp.admins_list'))

    admin = Admin(username=username, email=email)
    admin.set_password(password)
    db.session.add(admin)
    db.session.commit()
    flash(f'تم إضافة المشرف "{username}" بنجاح', 'success')
    return redirect(url_for('admin_bp.admins_list'))


@admin_bp.route('/admins/delete/<int:admin_id>', methods=['POST'])
@login_required
def admin_delete(admin_id):
    if admin_id == current_user.id:
        flash('لا يمكنك حذف حسابك الخاص', 'error')
        return redirect(url_for('admin_bp.admins_list'))
    admin = Admin.query.get_or_404(admin_id)
    db.session.delete(admin)
    db.session.commit()
    flash(f'تم حذف المشرف "{admin.username}"', 'success')
    return redirect(url_for('admin_bp.admins_list'))


# ─── Homepage Settings ───
@admin_bp.route('/homepage', methods=['GET', 'POST'])
@login_required
def homepage_settings():
    if request.method == 'POST':
        # Update featured products
        FeaturedProduct.query.delete()
        product_ids = request.form.getlist('featured_products', type=int)
        for i, pid in enumerate(product_ids):
            db.session.add(FeaturedProduct(product_id=pid, sort_order=i))

        # Update company info numbers
        for key in ['years_experience', 'clients_count', 'products_count', 'coverage']:
            val = request.form.get(key, '')
            info = CompanyInfo.query.filter_by(key=key).first()
            if info:
                info.value = val
            elif val:
                db.session.add(CompanyInfo(key=key, value=val))

        db.session.commit()
        flash('تم تحديث إعدادات الصفحة الرئيسية', 'success')
        return redirect(url_for('admin_bp.homepage_settings'))

    products = Product.query.filter_by(is_active=True).order_by(Product.name_ar).all()
    featured_ids = [f.product_id for f in FeaturedProduct.query.order_by(FeaturedProduct.sort_order).all()]
    info = {item.key: item.value for item in CompanyInfo.query.all()}

    return render_template('admin/homepage_settings.html', products=products, featured_ids=featured_ids, info=info)


# ─── Shipping Settings ───
@admin_bp.route('/shipping', methods=['GET', 'POST'])
@login_required
def shipping_settings():
    if request.method == 'POST':
        ship_id = request.form.get('id', type=int)
        if ship_id:
            ship = ShippingInfo.query.get(ship_id)
            ship.delivery_area_ar = request.form.get('delivery_area_ar', '')
            ship.estimated_days = request.form.get('estimated_days', '')
            ship.coverage_regions_ar = request.form.get('coverage_regions_ar', '')
        else:
            ship = ShippingInfo(
                delivery_area_ar=request.form.get('delivery_area_ar', ''),
                estimated_days=request.form.get('estimated_days', ''),
                coverage_regions_ar=request.form.get('coverage_regions_ar', '')
            )
            db.session.add(ship)
        db.session.commit()
        flash('تم تحديث معلومات الشحن', 'success')
        return redirect(url_for('admin_bp.shipping_settings'))

    shipping = ShippingInfo.query.all()
    return render_template('admin/shipping_settings.html', shipping=shipping)
